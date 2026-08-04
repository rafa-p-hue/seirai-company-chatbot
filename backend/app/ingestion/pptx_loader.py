"""PowerPoint parser — one or more semantic chunks per slide."""

from __future__ import annotations

import io
from typing import Any, Dict, List, Tuple

from pptx import Presentation

from app.ingestion.formats import EmptyDocumentError, PasswordProtectedError
from app.models.api import ExtractedPage, ExtractedTableRow


def extract_pptx_bytes(
    data: bytes, *, filename: str = "document.pptx"
) -> Tuple[List[ExtractedPage], Dict[str, Any]]:
    if not data:
        raise EmptyDocumentError("PPTX file is empty.")
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        message = str(exc).lower()
        if "password" in message or "encrypted" in message:
            raise PasswordProtectedError(
                "This PowerPoint file appears password-protected and cannot be processed."
            ) from exc
        raise ValueError(f"Could not parse PPTX file: {exc}") from exc

    pages: List[ExtractedPage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        title = _slide_title(slide)
        lines: List[str] = []
        table_rows: List[ExtractedTableRow] = []
        if title:
            lines.append(title)
        table_index = 0
        for shape in slide.shapes:
            if shape.has_table:
                rows, human = _shape_table(
                    shape.table, table_index=table_index, heading=title
                )
                table_rows.extend(rows)
                lines.extend(human)
                table_index += 1
                continue
            if not shape.has_text_frame:
                continue
            if title and _shape_text(shape).strip() == title:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if not text:
                    text = (paragraph.text or "").strip()
                if not text:
                    continue
                # Prefer bullets/sentences so body lines are not mistaken for headings.
                if text.startswith(("•", "-", "*")):
                    lines.append(
                        text if text.startswith("•") else f"• {text.lstrip('-* ')}"
                    )
                elif paragraph.level and paragraph.level > 0:
                    lines.append(f"• {text}")
                elif text.endswith((".", "!", "?")):
                    lines.append(text)
                else:
                    lines.append(f"• {text}")

        notes = _speaker_notes(slide)
        if notes:
            lines.append(f"Speaker notes: {notes}")

        content = "\n".join(line for line in lines if line).strip()
        if not content and not table_rows:
            continue
        pages.append(
            ExtractedPage(
                page_number=index,
                text=content or title or f"Slide {index}",
                table_rows=table_rows,
                slide_number=index,
                section_heading=title,
                metadata={"filename": filename, "slide_number": index, "title": title},
            )
        )

    if not pages:
        raise EmptyDocumentError(
            "No useful text could be extracted from this PowerPoint file."
        )
    return pages, {"page_count": len(pages), "slide_count": len(pages)}


def _slide_title(slide) -> str | None:
    if slide.shapes.title is not None:
        text = (slide.shapes.title.text or "").strip()
        if text:
            return text
    for shape in slide.shapes:
        if shape.has_text_frame and getattr(shape, "is_placeholder", False):
            try:
                if shape.placeholder_format.type is not None:
                    text = _shape_text(shape)
                    if text:
                        return text
            except Exception:  # noqa: BLE001
                continue
    return None


def _shape_text(shape) -> str:
    if not shape.has_text_frame:
        return ""
    return (shape.text or "").strip()


def _speaker_notes(slide) -> str:
    try:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            return (slide.notes_slide.notes_text_frame.text or "").strip()
    except Exception:  # noqa: BLE001
        return ""
    return ""


def _shape_table(table, *, table_index: int, heading: str | None):
    matrix = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
    if not matrix:
        return [], []
    headers = matrix[0]
    out: List[ExtractedTableRow] = []
    human: List[str] = []
    for row_index, cells in enumerate(matrix[1:] or [], start=1):
        while len(cells) < len(headers):
            cells.append("")
        cells = cells[: len(headers)]
        parts = [f"{headers[i]}: {cells[i]}" for i in range(len(headers)) if cells[i]]
        human_text = "; ".join(parts)
        out.append(
            ExtractedTableRow(
                table_index=table_index,
                row_index=row_index,
                heading=heading,
                column_labels=list(headers),
                cells=list(cells),
                row_label=cells[0] if cells else None,
                human_text=human_text,
            )
        )
        human.append(human_text)
    return out, human
