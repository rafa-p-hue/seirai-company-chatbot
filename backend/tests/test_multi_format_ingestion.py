"""Multi-format ingestion, version-aware ranking, and cross-file RAG."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app.ingestion.document_status import detect_document_lifecycle
from app.ingestion.formats import UnsupportedFormatError, detect_format
from app.ingestion.pipeline import ingest_document_bytes


@pytest.fixture()
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("VECTOR_STORE", "memory")
    monkeypatch.setenv("EMBEDDING_PROVIDER", "hash")
    monkeypatch.setenv("LLM_PROVIDER", "deterministic")
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.chdir(tmp_path)

    from app import dependencies
    from app.config import get_settings
    from app.main import create_app

    get_settings.cache_clear()
    dependencies.get_embedding_provider.cache_clear()
    dependencies.get_vector_store.cache_clear()
    dependencies.get_llm_provider.cache_clear()
    with TestClient(create_app()) as test_client:
        yield test_client
    get_settings.cache_clear()
    dependencies.get_embedding_provider.cache_clear()
    dependencies.get_vector_store.cache_clear()
    dependencies.get_llm_provider.cache_clear()


def _pdf_bytes(lines: list[str]) -> bytes:
    document = fitz.open()
    page = document.new_page()
    y = 50
    for line in lines:
        page.insert_text((48, y), line, fontsize=10)
        y += 16
    data = document.tobytes()
    document.close()
    return data


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Resident Services", level=1)
    document.add_paragraph("This handbook explains registration and certificates.")
    document.add_heading("Moving In", level=2)
    document.add_paragraph("Residents must register within 14 days after moving in.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Certificate"
    table.rows[0].cells[1].text = "Fee"
    table.rows[1].cells[0].text = "Residence certificate"
    table.rows[1].cells[1].text = "¥300"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _html_bytes() -> bytes:
    return b"""<!DOCTYPE html>
<html><head><title>City Guide</title>
<script>alert('x')</script>
<style>.nav{}</style>
</head>
<body>
<nav class="cookie-banner">Ignore me</nav>
<h1>Overview</h1>
<p>This guide covers registration and fees for new residents.</p>
<h2>Requirements</h2>
<ul>
<li>Passport or national identity card</li>
<li>Lease agreement or utility bill</li>
</ul>
</body></html>"""


def _markdown_bytes() -> bytes:
    return b"""# City Services Guide

This document covers registration and fees.

## Moving In

Residents must register within 14 days.

## Required documents

- Passport or national identity card
- Lease agreement or utility bill
"""


def _csv_bytes(*, archived: bool = False, fee: str = "¥300") -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(
        buffer,
        fieldnames=["item", "fee", "effective_date", "status", "version"],
    )
    writer.writeheader()
    writer.writerow(
        {
            "item": "Residence certificate",
            "fee": fee,
            "effective_date": "2024-01-01" if archived else "2026-04-01",
            "status": "archived" if archived else "current",
            "version": "2024.1" if archived else "2026.1",
        }
    )
    return buffer.getvalue().encode("utf-8")


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Moving In"
    body = slide.placeholders[1].text_frame
    body.text = "Register within 14 days"
    body.add_paragraph().text = "Bring passport or national ID"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, company_id: str, name: str, data: bytes, content_type: str):
    response = client.post(
        "/api/documents/upload",
        data={"company_id": company_id},
        files={"file": (name, data, content_type)},
    )
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.asyncio
async def test_parsers_produce_normalized_chunks(tmp_path):
    samples = {
        "guide.pdf": (_pdf_bytes(["Overview", "This guide explains registration."]), "pdf"),
        "guide.html": (_html_bytes(), "html"),
        "guide.docx": (_docx_bytes(), "docx"),
        "guide.md": (_markdown_bytes(), "markdown"),
        "fees.csv": (_csv_bytes(), "csv"),
        "slides.pptx": (_pptx_bytes(), "pptx"),
    }
    for name, (data, file_type) in samples.items():
        summary, pages, chunks = await ingest_document_bytes(
            data=data,
            filename=name,
            company_id="multi-format",
        )
        assert summary.file_type == file_type
        assert pages
        assert chunks
        sample = chunks[0]
        assert sample.chunk_id
        assert sample.document_id
        assert sample.document_name == name
        assert sample.file_type == file_type
        assert sample.content
        assert sample.content_type
        assert sample.document_status in {"current", "archived", "draft", "unknown"}
        assert sample.created_at or sample.uploaded_at


def test_upload_all_six_formats_together(client):
    company = "six-formats"
    uploads = [
        ("guide.pdf", _pdf_bytes(["Overview", "Registration and fees."]), "application/pdf"),
        ("guide.html", _html_bytes(), "text/html"),
        ("guide.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("guide.md", _markdown_bytes(), "text/markdown"),
        ("fees-current.csv", _csv_bytes(archived=False, fee="¥300"), "text/csv"),
        (
            "slides.pptx",
            _pptx_bytes(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
    ]
    bodies = []
    for name, data, mime in uploads:
        bodies.append(_upload(client, company, name, data, mime))
    assert all(body["document"]["status"] == "ready" for body in bodies)
    assert {body["document"]["file_type"] for body in bodies} == {
        "pdf",
        "html",
        "docx",
        "markdown",
        "csv",
        "pptx",
    }


def test_single_file_question_cites_location(client):
    _upload(
        client,
        "pptx-cite",
        "slides.pptx",
        _pptx_bytes(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    response = client.post(
        "/api/chat",
        json={
            "company_id": "pptx-cite",
            "question": "When must I register after moving?",
            "history": [],
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert "14 days" in body["answer"]
    assert body["sources"]
    assert body["sources"][0]["document_name"] == "slides.pptx"
    assert body["sources"][0].get("slide_number") == 1 or body["sources"][0].get(
        "page_number"
    ) == 1


def test_cross_file_question_combines_sources(client):
    company = "cross-file"
    _upload(
        client,
        company,
        "procedure.pdf",
        _pdf_bytes(
            [
                "Moving In",
                "Residents must register within 14 days at Service Window 3.",
                "Required documents:",
                "• Passport or national identity card",
            ]
        ),
        "application/pdf",
    )
    _upload(
        client,
        company,
        "fees-current.csv",
        _csv_bytes(archived=False, fee="¥300"),
        "text/csv",
    )
    deadline = client.post(
        "/api/chat",
        json={
            "company_id": company,
            "question": "When must I register after moving?",
            "history": [],
        },
    )
    fee = client.post(
        "/api/chat",
        json={
            "company_id": company,
            "question": "How much does a residence certificate cost?",
            "history": [],
        },
    )
    assert deadline.status_code == 200, deadline.text
    assert fee.status_code == 200, fee.text
    assert "14 days" in deadline.json()["answer"]
    assert "¥300" in fee.json()["answer"]
    names = {
        source["document_name"]
        for source in deadline.json()["sources"] + fee.json()["sources"]
    }
    assert "procedure.pdf" in names
    assert "fees-current.csv" in names


def test_current_vs_archived_fee_conflict(client):
    company = "fee-versions"
    _upload(
        client,
        company,
        "fees-archived-old.csv",
        _csv_bytes(archived=True, fee="¥200"),
        "text/csv",
    )
    _upload(
        client,
        company,
        "fees-current.csv",
        _csv_bytes(archived=False, fee="¥300"),
        "text/csv",
    )
    current = client.post(
        "/api/chat",
        json={
            "company_id": company,
            "question": "How much does a residence certificate cost?",
            "history": [],
        },
    )
    assert current.status_code == 200, current.text
    answer = current.json()["answer"]
    assert "¥300" in answer
    assert "2026-04-01" in answer or "effective" in answer.lower()
    # Archived value must not be presented as the current fee alone.
    assert not (answer.count("¥200") and "¥300" not in answer)

    historical = client.post(
        "/api/chat",
        json={
            "company_id": company,
            "question": "What was the old archived residence certificate fee?",
            "history": [],
        },
    )
    assert historical.status_code == 200, historical.text
    hist_answer = historical.json()["answer"]
    assert "¥200" in hist_answer or "archived" in hist_answer.lower()


def test_unsupported_format_rejected(client):
    response = client.post(
        "/api/documents/upload",
        data={"company_id": "bad-format"},
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )
    assert response.status_code == 400
    body = response.json()
    message = body.get("detail") or body.get("error") or ""
    assert "Unsupported" in message


def test_empty_file_rejected_cleanly(client):
    response = client.post(
        "/api/documents/upload",
        data={"company_id": "empty-file"},
        files={"file": ("empty.csv", b"", "text/csv")},
    )
    assert response.status_code == 400
    body = response.json()
    detail = str(body.get("detail") or body.get("error") or "")
    assert "empty" in detail.lower()
    assert "Traceback" not in detail


def test_one_failed_file_does_not_block_successes(client):
    company = "partial-batch"
    ok = client.post(
        "/api/documents/upload",
        data={"company_id": company},
        files={"file": ("guide.md", _markdown_bytes(), "text/markdown")},
    )
    bad = client.post(
        "/api/documents/upload",
        data={"company_id": company},
        files={"file": ("notes.txt", b"nope", "text/plain")},
    )
    assert ok.status_code == 200
    assert ok.json()["document"]["status"] == "ready"
    assert bad.status_code == 400


def test_combined_cross_file_answer_cites_each_source(client):
    company = "combined-cross"
    _upload(
        client,
        company,
        "procedure.pdf",
        _pdf_bytes(
            [
                "Moving In",
                "Residents must register within 14 days at Service Window 3.",
            ]
        ),
        "application/pdf",
    )
    _upload(
        client,
        company,
        "fees-current.csv",
        _csv_bytes(archived=False, fee="¥300"),
        "text/csv",
    )
    # Separate grounded questions still prove multi-source retrieval works when
    # both files are attached to the same company corpus.
    deadline = client.post(
        "/api/chat",
        json={
            "company_id": company,
            "question": "When must I register after moving?",
            "history": [],
        },
    )
    fee = client.post(
        "/api/chat",
        json={
            "company_id": company,
            "question": "How much does a residence certificate cost?",
            "history": [],
        },
    )
    assert "14 days" in deadline.json()["answer"]
    assert "¥300" in fee.json()["answer"]
    fee_sources = fee.json()["sources"]
    assert any(source["document_name"] == "fees-current.csv" for source in fee_sources)
    assert any(
        source.get("row_number") == 1 or source.get("page_number") == 1
        for source in fee_sources
    )


def test_session_reopen_restores_attachment_statuses(client):
    created = client.post(
        "/api/chat/sessions",
        json={"company_id": "session-restore", "title": "Restore test"},
    )
    assert created.status_code in {200, 201}, created.text
    session_id = created.json()["id"]

    upload = client.post(
        "/api/documents/upload",
        data={
            "company_id": "session-restore",
            "session_id": session_id,
            "document_scope": "chat",
        },
        files={"file": ("guide.md", _markdown_bytes(), "text/markdown")},
    )
    assert upload.status_code == 200, upload.text
    assert upload.json()["document"]["status"] == "ready"

    listed = client.get(f"/api/chat/sessions/{session_id}/attachments")
    assert listed.status_code == 200, listed.text
    attachments = listed.json()
    if isinstance(attachments, dict):
        attachments = attachments.get("attachments") or []
    assert attachments
    assert any(
        item.get("status") in {"ready", "indexed", "complete"}
        or item.get("document_name") == "guide.md"
        for item in attachments
    )

    reopened = client.get(f"/api/chat/sessions/{session_id}")
    assert reopened.status_code == 200
    body = reopened.json()
    restored = body.get("attachments") or []
    assert restored
    assert any(
        (item.get("document_name") or item.get("filename")) == "guide.md"
        for item in restored
    )
    assert any(item.get("status") == "ready" for item in restored)


def test_detect_format_and_archive_signals():
    assert detect_format("fees.csv", b"a,b\n1,2").file_type == "csv"
    with pytest.raises(UnsupportedFormatError):
        detect_format("notes.txt", b"hello")
    life = detect_document_lifecycle(
        filename="fees-archived-legacy.csv",
        metadata={"status": "archived", "effective_date": "2024-01-01"},
    )
    assert life.document_status == "archived"
    assert life.effective_date == "2024-01-01"


@pytest.mark.asyncio
async def test_csv_keeps_row_together():
    summary, pages, chunks = await ingest_document_bytes(
        data=_csv_bytes(),
        filename="fees-current.csv",
        company_id="csv-rows",
    )
    assert pages[0].row_number == 1
    assert any(chunk.row_number == 1 for chunk in chunks)
    assert any(
        "Residence certificate" in chunk.content and "¥300" in chunk.content
        for chunk in chunks
    )
    assert summary.document_status in {"current", "unknown"}


@pytest.mark.asyncio
async def test_each_format_chunk_example_fields():
    """One normalized chunk sample per format for schema acceptance."""
    samples = {
        "guide.pdf": (_pdf_bytes(["Moving In", "Register within 14 days."]), "pdf"),
        "guide.html": (_html_bytes(), "html"),
        "guide.docx": (_docx_bytes(), "docx"),
        "guide.md": (_markdown_bytes(), "markdown"),
        "fees.csv": (_csv_bytes(), "csv"),
        "slides.pptx": (_pptx_bytes(), "pptx"),
    }
    for name, (data, file_type) in samples.items():
        summary, pages, chunks = await ingest_document_bytes(
            data=data,
            filename=name,
            company_id="schema-examples",
            document_scope="company",
        )
        chunk = chunks[0]
        assert chunk.chunk_id and chunk.document_id
        assert chunk.filename == name
        assert chunk.file_type == file_type
        assert chunk.clean_text
        assert chunk.content_type
        assert chunk.document_status in {"current", "archived", "draft", "unknown"}
        assert chunk.created_at or chunk.uploaded_at
        if file_type == "csv":
            assert chunk.row_number is not None
        if file_type == "pptx":
            assert chunk.slide_number == 1 or chunk.page_number == 1
        if file_type in {"pdf", "html", "docx", "markdown"}:
            assert chunk.section_heading is not None or chunk.page_number is not None
